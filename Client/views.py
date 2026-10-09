from django.views.generic.edit import CreateView
from django.views.generic import ListView, UpdateView, DeleteView
from .models import Client
from django.urls import reverse_lazy
from .forms import ClientCreationForm
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
#Для DeleteView
from django.db.models import ProtectedError
from django.shortcuts import redirect

#для excel и Pdf
from django.conf import settings
from django.http import HttpResponse
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

#Для поля поиска
from django.db.models import Prefetch
from django.http import JsonResponse
from adress.models import ClientAddress
from django.db import transaction
from django.http import HttpResponseRedirect
from django.utils.html import format_html

from django.db.models import Max
from django.shortcuts import render
from django.core.paginator import Paginator


class AddressTakenError(Exception):
    pass


class ClientSaveMixin:
    def form_valid(self, form):
        # Как и раньше, автором становится текущий пользователь
        form.instance.author = self.request.user
        new_ids = {a.pk for a in form.cleaned_data['addresses']}

        try:
            with transaction.atomic():
                self.object = form.save()
                current_ids = set(self.object.addresses.values_list('pk', flat=True))
                to_remove = current_ids - new_ids
                to_add = new_ids - current_ids

                if to_remove:
                    ClientAddress.objects.filter(
                        pk__in=to_remove, client=self.object
                    ).update(client=None)

                if to_add:
                    updated = ClientAddress.objects.filter(
                        pk__in=to_add, client__isnull=True
                    ).update(client=self.object)
                    if updated != len(to_add):
                        raise AddressTakenError
        except AddressTakenError:
            form.add_error('addresses', 'Деякі адреси вже зайняті іншим клієнтом. Оберіть інші.')
            return self.form_invalid(form)

        messages.success(
            self.request,
            format_html(self.success_message, self.object.client_name),
        )
        return HttpResponseRedirect(self.get_success_url())


class ClientCreateView(ClientSaveMixin, LoginRequiredMixin, CreateView):
    model = Client
    template_name = 'client_new.html'
    form_class = ClientCreationForm
    success_url = reverse_lazy('home')
    success_message = 'Клієнт <strong>{}</strong> вдало створений!'


class ClientUpdateView(ClientSaveMixin, LoginRequiredMixin, UpdateView):
    model = Client
    template_name = 'client_edit.html'
    form_class = ClientCreationForm
    success_url = reverse_lazy('client_list')
    success_message = 'Запис <strong>{}</strong> було змінено!'


class ClientListView(LoginRequiredMixin, ListView):
  model = Client
  template_name = 'client_list.html'
  context_object_name = "clients"
  paginate_by = 10
  ordering = ["-created_at"]

  # Поиск в фильтре
  def get_queryset(self):
    queryset = super().get_queryset().prefetch_related('addresses')
    q = self.request.GET.get("q", "").strip()
    if q:
      queryset = queryset.filter(client_code__icontains=q)
    return queryset

  # Сохраняет искомое значение при переходах по страницам пагинации
  # def get_context_data(self, **kwargs):
  #   context = super().get_context_data(**kwargs)
  #   context["q"] = self.request.GET.get("q", "").strip()
  #   return context

  def get_context_data(self, **kwargs):
      context = super().get_context_data(**kwargs)
      context["q"] = self.request.GET.get("q", "").strip()
      # курсор для опроса: самый большой pk среди ВСЕХ клиентов
      context["last_id"] = Client.objects.aggregate(m=Max("pk"))["m"] or 0
      return context


class ClientNewItemsView(LoginRequiredMixin, View):
    def get(self, request):
        try:
            after = int(request.GET.get("after", 0))
        except ValueError:
            after = 0
        q = request.GET.get("q", "").strip()

        newer = Client.objects.filter(pk__gt=after)
        last_id = newer.aggregate(m=Max("pk"))["m"]
        if last_id is None:
            return HttpResponse(status=204)

        clients = newer.prefetch_related("addresses").order_by("-pk")
        all_clients = Client.objects.order_by("-created_at")
        if q:
            clients = clients.filter(client_code__icontains=q)
            all_clients = all_clients.filter(client_code__icontains=q)

        # пагинация первой страницы (опрос работает только на ней)
        paginator = Paginator(all_clients, ClientListView.paginate_by)
        page_obj = paginator.get_page(1)

        return render(request, "client_new_items.html", {
            "clients": clients,
            "last_id": last_id,
            "q": q,
            "paginator": paginator,
            "page_obj": page_obj,
            "is_paginated": paginator.num_pages > 1,
        })


class ClientDeleteView(LoginRequiredMixin, DeleteView):
    model = Client
    success_url = reverse_lazy("client_list")
    context_object_name = "client"

    # Теперь это не страница, а фрагмент, который htmx вставит в модальное окно
    template_name = "client_delete_modal.html"

    def get(self, request, *args, **kwargs):
      # htmx добавляет к каждому запросу заголовок HX-Request.
      # Если кто-то откроет адрес в браузере напрямую, не показываем
      # голый фрагмент, а возвращаем на список
      if not request.headers.get("HX-Request"):
        return redirect("client_list")
      return super().get(request, *args, **kwargs)

    def form_valid(self, form):
      # Без изменений: удаление и сообщение (POST)
      name = self.object.client_name
      try:
        response = super().form_valid(form)
      except ProtectedError:
        messages.error(
          self.request,
          f"Неможливо видалити <strong>«{name}»</strong>: у клієнта є пов'язані адреси.",
        )
        return redirect("client_list")
      messages.success(self.request, f"Клієнта <strong>«{name}»</strong> видалено.")
      return response
#>>>>>>>>>>>>>>>>>>>>>>>> PDF и Excel
from django.shortcuts import get_object_or_404

EXPORT_FIELDS = [
    'client_code',
    'client_name',
    'client_registration_region',
    'client_registration_city',
    'email',
    'client_phone_number',
]


class ClientSingleExportMixin(LoginRequiredMixin):
    def get_client(self):
        # если клиенты привязаны к пользователю, добавьте author=self.request.user
        return get_object_or_404(Client, pk=self.kwargs['pk'])

    @staticmethod
    def get_rows(client):
        rows = []
        for f in EXPORT_FIELDS:
            label = str(Client._meta.get_field(f).verbose_name)
            value = getattr(client, f)
            rows.append((label, '' if value is None else str(value)))
        return rows

class ClientExcelExportView(ClientSingleExportMixin, View):
    def get(self, request, *args, **kwargs):
        client = self.get_client()

        wb = Workbook()
        ws = wb.active
        ws.title = 'Клієнт'

        ws.append(['Поле', 'Значення'])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for label, value in self.get_rows(client):
            ws.append([label, value])

        for idx, column in enumerate(ws.columns, start=1):
            max_len = max(len(str(c.value or '')) for c in column)
            ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 60)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="client_{client.pk}.xlsx"'
        wb.save(response)
        return response

class ClientPdfExportView(ClientSingleExportMixin, View):
    FONT_NAME = 'DejaVuSans'

    def get(self, request, *args, **kwargs):
        client = self.get_client()

        if self.FONT_NAME not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(
                TTFont(self.FONT_NAME, settings.BASE_DIR / 'fonts' / 'DejaVuSans.ttf')
            )

        style = ParagraphStyle('cell', fontName=self.FONT_NAME, fontSize=10, leading=13)
        head_style = ParagraphStyle('head', parent=style, textColor=colors.white)

        data = [[Paragraph('Поле', head_style), Paragraph('Значення', head_style)]]
        for label, value in self.get_rows(client):
            data.append([Paragraph(label, style), Paragraph(value, style)])

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="client_{client.pk}.pdf"'

        doc = SimpleDocTemplate(response, pagesize=A4,
                                leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
        table = Table(data, colWidths=[170, None], repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')]),
        ]))
        doc.build([table])
        return response